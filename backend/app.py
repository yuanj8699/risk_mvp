"""本地业务 API：导入、规则监测、证据保存、函件审核与模拟下发。

Pi 只负责选择辅助指标和分析证据；所有数据库写入及函件状态迁移由此处控制。
启动：python -m uvicorn backend.app:app --host 127.0.0.1 --port 8000
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from contextlib import asynccontextmanager
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, TypeVar

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import func, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .analysis import AnalysisError, provider
from .models import (ROOT, Base, engine, SessionLocal, Organization, Indicator, Snapshot, Observation,
                     Rule, AnalysisRun, RiskCase, Letter, LetterHistory, AuditLog, uid, now)
from .seed import seed_database
from .services import (
    number, display, observation_issue, evaluate, evidence_record,
    search_indicators, query_metrics, check_risks, get_letter_template,
    freeze_context, load_task_context,
)

ModelType = TypeVar("ModelType", bound=Base)


def hash_json(value: Any) -> str:
    """按稳定键序计算摘要，用于运行去重及函件内容绑定。"""
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def timestamp(value: datetime) -> str:
    # SQLite读取的时间可能不带时区；数据库中的时间统一按UTC写入。
    return value.replace(tzinfo=timezone.utc).isoformat() if value.tzinfo is None else value.isoformat()


# 序列化函数限定公开字段，避免直接暴露ORM对象及内部状态。
def org_dict(row: Organization) -> dict:
    return {k: getattr(row, k) for k in ("id", "code", "name", "level", "parent_id", "region_id")}


def metric_dict(row: Indicator) -> dict:
    return {k: getattr(row, k) for k in ("id", "name", "unit", "frequency", "description")}


def rule_dict(row: Rule) -> dict:
    return {**{k: getattr(row, k) for k in ("id", "name", "kind", "indicator_id", "enabled", "params", "version", "description")},
            "needs_config": row.kind == "ratio" and row.params.get("threshold") is None}


def snapshot_dict(row: Snapshot) -> dict:
    return {"id": row.id, "name": row.name, "data_date": row.data_date.isoformat(), "source": row.source,
            "is_demo": row.is_demo, "row_count": row.row_count, "status": row.status}


def run_dict(row: AnalysisRun) -> dict:
    return {"id": row.id, "status": row.status, "snapshot_id": row.snapshot_id, "created_at": timestamp(row.created_at),
            "summary": row.summary, "selection": row.selection}


def case_dict(row: RiskCase) -> dict:
    return {"id": row.id, "run_id": row.run_id, "org_id": row.org_id, "org_name": row.org_name,
            "region_name": row.region_name, "data_date": row.data_date.isoformat(), "risk_types": row.risk_types,
            "status": row.status, "evidence_count": len(row.evidence), "headline": row.headline}


def letter_dict(row: Letter) -> dict:
    return {"id": row.id, "case_id": row.case_id, "revision": row.revision, "content": row.content,
            "status": row.status, "evidence_hash": row.evidence_hash, "created_at": timestamp(row.created_at), "receipt": row.receipt}


def get_or_404(db: Session, model: type[ModelType], key: str) -> ModelType:
    """供各API复用的主键查询；不存在时统一返回404。"""
    value = db.get(model, key)
    if value is None:
        raise HTTPException(404, "记录不存在")
    return value


def audit(db: Session, action: str, key: str, details: dict | None = None) -> None:
    """操作记录随调用方事务提交，不能脱离业务写入单独成功。"""
    db.add(AuditLog(action=action, entity_id=key, details=details or {}))


@asynccontextmanager
async def lifespan(app):
    """首次启动创建测试结构和样例；已有业务数据不会被种子覆盖。"""
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        seed_database(db)
    yield


app = FastAPI(title="行社风险提示测试版", version="0.1.0", lifespan=lifespan)


@app.exception_handler(AnalysisError)
async def analysis_error_handler(request, exc):
    return JSONResponse(status_code=502, content={"detail": str(exc)})


class Payload(BaseModel):
    model_config = ConfigDict(extra="forbid")


class RunPayload(Payload):
    snapshot_id: str
    risk_topic: str = Field(default="", max_length=500)
    org_scope: list[str] | None = None
    retry: bool = False


class RulePayload(Payload):
    enabled: bool | None = None
    params: dict | None = None


class RevisionPayload(Payload):
    """调用方看到的版本号，用于拒绝覆盖其他人的新版本。"""
    expected_revision: int = Field(ge=1)


class EditPayload(RevisionPayload):
    content: str = Field(min_length=1, max_length=30000)


@app.get("/api/health")
def health():
    return {"status": "ok", "demo": True, "real_delivery": False}


@app.get("/api/bootstrap")
def bootstrap():
    """一次取得工作台初始化所需的目录、规则、快照和最近运行。"""
    with SessionLocal() as db:
        orgs = db.scalars(select(Organization).order_by(Organization.level, Organization.code)).all()
        indicators = db.scalars(select(Indicator).order_by(Indicator.id)).all()
        rules = db.scalars(select(Rule).order_by(Rule.id)).all()
        snaps = db.scalars(select(Snapshot).order_by(Snapshot.created_at.desc(), Snapshot.data_date.desc())).all()
        latest = db.scalar(select(AnalysisRun).order_by(AnalysisRun.created_at.desc()).limit(1))
        return {
            "meta": {"app_name": "行社风险提示", "version": "0.1.0", "demo": True, "agent_mode": provider.mode,
                     "data_date": max((s.data_date for s in snaps), default=date.today()).isoformat(),
                     "rank_note": "默认排名≥72；82家全部参评且无并列时为第72—82名，共11家；并列不截断。"},
            "summary": {"region_count": sum(o.level == 2 for o in orgs), "org_count": sum(o.level == 3 for o in orgs),
                        "indicator_count": len(indicators), "enabled_rule_count": sum(r.enabled for r in rules),
                        "case_count": db.scalar(select(func.count()).select_from(RiskCase)) or 0,
                        "letter_count": db.scalar(select(func.count()).select_from(Letter)) or 0},
            "organizations": [org_dict(o) for o in orgs], "indicators": [metric_dict(i) for i in indicators],
            "rules": [rule_dict(r) for r in rules], "snapshots": [snapshot_dict(s) for s in snaps],
            "latest_run": run_dict(latest) if latest else None,
        }


@app.get("/api/organizations")
def organizations():
    with SessionLocal() as db:
        return [org_dict(o) for o in db.scalars(select(Organization).order_by(Organization.level, Organization.code))]


@app.get("/api/indicators")
def indicators(q: str = ""):
    return search_indicators(q, limit=1000)["items"]


@app.get("/api/runs/{run_id}/metrics")
def task_metrics(run_id: str, indicator_ids: str, org_id: str | None = None,
                 history: bool = False, limit: int = 100, offset: int = 0):
    try:
        context = load_task_context(run_id, org_id)
        return query_metrics(context, indicator_ids.split(","),
                             org_ids=[org_id] if org_id else None,
                             history=history, limit=limit, offset=offset)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@app.get("/api/rules")
def rules():
    with SessionLocal() as db:
        return [rule_dict(r) for r in db.scalars(select(Rule).order_by(Rule.id))]


@app.patch("/api/rules/{rule_id}")
def patch_rule(rule_id: str, data: RulePayload):
    """仅允许修改各类规则开放的参数；变动后递增版本，不改历史证据。"""
    with SessionLocal() as db:
        rule = get_or_404(db, Rule, rule_id)
        params = dict(rule.params)
        changes = data.params or {}
        allowed = {"rank": {"cutoff"}, "ratio": {"threshold", "op"}, "exists": set()}[rule.kind]
        if set(changes) - allowed:
            raise HTTPException(422, "该规则含不可修改参数")
        params.update(changes)
        if rule.kind == "rank":
            cutoff = params.get("cutoff")
            if isinstance(cutoff, bool) or not isinstance(cutoff, int) or not 1 <= cutoff <= 82:
                raise HTTPException(422, "排名边界必须为1至82之间的整数")
        elif rule.kind == "ratio":
            threshold = params.get("threshold")
            if threshold is not None:
                try:
                    value = Decimal(str(threshold))
                except InvalidOperation:
                    raise HTTPException(422, "占比阈值必须是有限数值")
                if not value.is_finite() or value < 0 or value > 100:
                    raise HTTPException(422, "占比阈值必须在0至100之间，单位为%")
                params["threshold"] = number(value)
            if params.get("op") not in {">=", ">"}:
                raise HTTPException(422, "占比比较符只支持>=或>")
        enabled = rule.enabled if data.enabled is None else data.enabled
        if params != rule.params or enabled != rule.enabled:
            rule.params, rule.enabled, rule.version = params, enabled, rule.version + 1
            audit(db, "rule_updated", rule.id, {"version": rule.version, "params": params, "test_configuration": True})
            db.commit()
        return rule_dict(rule)


@app.post("/api/runs")
def create_run(data: RunPayload):
    """先提交任务与检测证据，再在事务外执行Pi；失败保留记录，显式重试。"""
    with SessionLocal() as db:
        snapshot = get_or_404(db, Snapshot, data.snapshot_id)
        if snapshot.status != "ready":
            raise HTTPException(409, "数据快照尚未就绪")
        available_orgs = set(db.scalars(select(Organization.id).where(Organization.level == 3)))
        scope = sorted(set(data.org_scope) if data.org_scope is not None else available_orgs)
        if not scope or not set(scope) <= available_orgs:
            raise HTTPException(422, "监测机构范围必须为目录内的行社")
        frozen_rules = [rule_dict(r) for r in db.scalars(select(Rule).order_by(Rule.id))]
        topic = " ".join(data.risk_topic.split())
        key = hash_json({"snapshot": snapshot.id, "rules": frozen_rules, "risk_topic": topic,
                         "org_scope": scope, "agent": provider.cache_identity})
        run = db.scalar(select(AnalysisRun).where(AnalysisRun.run_key == key))
        if run:
            if run.status != "failed" or not data.retry:
                return run_dict(run)
            # 条件更新只允许一个重试者取走failed状态，其他请求复用该任务。
            claimed = db.execute(update(AnalysisRun).where(
                AnalysisRun.id == run.id, AnalysisRun.status == "failed",
            ).values(status="running"))
            if claimed.rowcount != 1:
                db.rollback()
                db.refresh(run)
                return run_dict(run)
            context = run.selection["task_context"]
            run.selection = {**run.selection, "error": None,
                             "attempt": run.selection.get("attempt", 1) + 1}
        else:
            run_id = uid()
            context = freeze_context(db, run_id, snapshot, topic, frozen_rules, scope)
            run = AnalysisRun(id=run_id, run_key=key, snapshot_id=snapshot.id, status="running",
                              summary={}, rule_snapshot=frozen_rules,
                              selection={"mode": provider.mode, "task_context": context,
                                         "indicator_ids": [], "attempt": 1})
            db.add(run)
        try:
            db.commit()
        except IntegrityError:
            db.rollback()
            winner = db.scalar(select(AnalysisRun).where(AnalysisRun.run_key == key))
            if winner:
                return run_dict(winner)
            raise
        run_id = run.id
        is_demo = snapshot.is_demo

    # 此处无存活的写事务；MCP使用另一个连接也能读取任务及固定上下文。
    try:
        detection = check_risks(context)
        with SessionLocal() as db:
            run = db.get(AnalysisRun, run_id)
            run.summary, run.checks = detection["summary"], detection["checks"]
            for item in detection["cases"]:
                case = db.scalar(select(RiskCase).where(
                    RiskCase.run_id == run_id, RiskCase.org_id == item["case_key"],
                ))
                if case is None:
                    types = sorted({{"rank": "排名提醒", "ratio": "占比提醒", "exists": "存在提醒"}[e["rule_kind"]]
                                    for e in item["evidence"]})
                    case = RiskCase(run_id=run_id, org_id=item["case_key"], org_name=item["org_name"],
                                    region_name=item["region_name"], data_date=date.fromisoformat(context["as_of"]),
                                    risk_types=types, status="open", evidence=item["evidence"],
                                    headline=f"{len(item['evidence'])}项规则触发 · " + "、".join(types))
                    db.add(case)
                case.analysis = {"mode": provider.mode, "status": "pending", "facts": [],
                                 "hypotheses": [], "checks": [], "limitations": ["分析尚未完成，暂不可拟函。"]}
            db.commit()
        catalog = search_indicators(limit=1000)["items"]
        mandatory = [r["indicator_id"] for r in context["rule_snapshot"] if r["enabled"]]
        selection = provider.select(context["topic"], catalog, mandatory)
        with SessionLocal() as db:
            run = db.get(AnalysisRun, run_id)
            run.selection = {**run.selection, **selection}
            db.commit()
        cases = [{"case_key": item["case_key"], "evidence": item["evidence"]}
                 for item in detection["cases"]]
        analyses = provider.explain_batch(cases, context)
        with SessionLocal() as db:
            run = db.get(AnalysisRun, run_id)
            for case in db.scalars(select(RiskCase).where(RiskCase.run_id == run_id)):
                analysis = analyses[case.org_id]
                if not is_demo:
                    analysis["limitations"] = [x for x in analysis["limitations"] if "内置数值" not in x]
                    analysis["limitations"].insert(0, "本批次为用户导入数据，未核验生产口径。")
                case.analysis = {**analysis, "status": "completed"}
            run.status = "completed"
            audit(db, "analysis_completed", run_id, {"snapshot_id": context["snapshot_id"]})
            db.commit()
            return run_dict(run)
    except Exception as exc:
        # 在业务边界保存失败状态；不把供应商响应、连接串等原始异常显示给用户。
        message = str(exc) if isinstance(exc, (AnalysisError, ValueError)) else "业务执行失败，请检查服务日志。"
        with SessionLocal() as db:
            run = db.get(AnalysisRun, run_id)
            run.status = "failed"
            run.selection = {**run.selection, "error": message}
            audit(db, "analysis_failed", run_id, {"error_type": type(exc).__name__})
            db.commit()
        raise HTTPException(502, {"run_id": run_id, "message": message}) from exc


@app.get("/api/runs")
def list_runs():
    with SessionLocal() as db:
        return [run_dict(r) for r in db.scalars(select(AnalysisRun).order_by(AnalysisRun.created_at.desc()))]


@app.get("/api/runs/{run_id}/checks")
def get_checks(run_id: str):
    with SessionLocal() as db:
        return get_or_404(db, AnalysisRun, run_id).checks


@app.get("/api/cases")
def list_cases(run_id: str | None = None):
    with SessionLocal() as db:
        query = select(RiskCase)
        if run_id:
            query = query.where(RiskCase.run_id == run_id)
        return [case_dict(c) for c in db.scalars(query.order_by(RiskCase.org_id))]


@app.get("/api/cases/{case_id}")
def get_case(case_id: str):
    with SessionLocal() as db:
        case = get_or_404(db, RiskCase, case_id)
        letter = db.scalar(select(Letter).where(Letter.case_id == case_id))
        return {**case_dict(case), "evidence": case.evidence, "analysis": case.analysis, "letter": letter_dict(letter) if letter else None}


def render_letter(case: RiskCase) -> str:
    """触发事实和补查事实都取自留存证据，模型只能贡献有引用的核查建议。"""
    facts = "\n".join(f"{i}. {e['indicator_name']}：{e['value_display']}，数据日期{e['data_date']}。{e['reason']}。"
                      for i, e in enumerate(case.evidence, 1))
    checks = "\n".join(f"{i}. {text}" for i, text in enumerate(case.analysis.get("checks", []), 1))
    refs = case.analysis.get("evidence_references", {})
    cited = {key for items in refs.values() for item in items for key in item.get("evidence_ids", [])}
    queried = [e for e in case.analysis.get("queried_evidence", []) if e["evidence_id"] in cited]
    supplementary = ""
    if queried:
        supplementary = "\n补充查询证据：\n" + "\n".join(
            f"{e['indicator_name']}：{e['value_display']}，数据日期{e['data_date']}，证据{e['evidence_id']}。"
            for e in queried) + "\n"
    return get_letter_template()["template"].format(
        org_name=case.org_name, data_date=case.data_date.isoformat(),
        facts=facts, checks=checks, supplementary=supplementary,
    )


@app.post("/api/cases/{case_id}/draft")
def draft_letter(case_id: str):
    """每个事项只建立一份函件；重复请求返回已有草稿。"""
    with SessionLocal() as db:
        case = get_or_404(db, RiskCase, case_id)
        existing = db.scalar(select(Letter).where(Letter.case_id == case_id))
        if existing:
            return letter_dict(existing)
        if get_or_404(db, AnalysisRun, case.run_id).status != "completed":
            raise HTTPException(409, "任务分析尚未完成，不能生成函件")
        evidence_hash = hash_json({"triggered": case.evidence, "queried": case.analysis.get("queried_evidence", [])})
        letter = Letter(case_id=case.id, content=render_letter(case), evidence_hash=evidence_hash, revision=1, status="draft")
        db.add(letter)
        try:
            db.flush()
            db.add(LetterHistory(letter_id=letter.id, revision=1, content=letter.content, evidence_hash=letter.evidence_hash))
            case.status = "draft"
            audit(db, "letter_drafted", letter.id, {"revision": 1})
            db.commit()
        except IntegrityError:
            db.rollback()
            letter = db.scalar(select(Letter).where(Letter.case_id == case_id))
            if not letter:
                raise HTTPException(409, "函件正在创建，请稍后重试")
        return letter_dict(letter)


def ensure_revision(letter: Letter, revision: int) -> None:
    if letter.revision != revision:
        raise HTTPException(409, "函件版本已变化，请刷新后重试")


def letter_content_hash(letter: Letter) -> str:
    """审批与发送共用绑定口径：正文、证据摘要和收函事项。"""
    return hash_json({
        "content": letter.content,
        "evidence": letter.evidence_hash,
        "case_id": letter.case_id,
    })


@app.patch("/api/letters/{letter_id}")
def edit_letter(letter_id: str, data: EditPayload):
    """编辑产生新修订并使原审批失效；条件UPDATE防止并发覆盖。"""
    with SessionLocal() as db:
        letter = get_or_404(db, Letter, letter_id)
        ensure_revision(letter, data.expected_revision)
        if letter.status == "simulated_sent":
            raise HTTPException(409, "已模拟下发的版本不可修改")
        content = data.content.strip()
        if not content:
            raise HTTPException(422, "函件内容不能为空")
        result = db.execute(update(Letter).where(Letter.id == letter_id, Letter.revision == data.expected_revision, Letter.status != "simulated_sent")
                            .values(content=content, revision=data.expected_revision + 1, status="draft", approved_revision=None, approved_hash=None))
        if result.rowcount != 1:
            raise HTTPException(409, "函件已被其他操作修改")
        db.add(LetterHistory(letter_id=letter_id, revision=data.expected_revision + 1, content=content, evidence_hash=letter.evidence_hash))
        get_or_404(db, RiskCase, letter.case_id).status = "draft"
        audit(db, "letter_edited", letter_id, {"revision": data.expected_revision + 1, "approval_invalidated": True})
        db.commit()
        db.refresh(letter)
        return letter_dict(letter)


@app.post("/api/letters/{letter_id}/approve")
def approve_letter(letter_id: str, data: RevisionPayload):
    """审批绑定当前版本和内容摘要，不自动覆盖后续修改。"""
    with SessionLocal() as db:
        letter = get_or_404(db, Letter, letter_id)
        ensure_revision(letter, data.expected_revision)
        if letter.status == "simulated_sent":
            return letter_dict(letter)
        result = db.execute(update(Letter).where(Letter.id == letter_id, Letter.revision == data.expected_revision, Letter.status.in_(["draft", "approved"]))
                            .values(status="approved", approved_revision=letter.revision, approved_hash=letter_content_hash(letter)))
        if result.rowcount != 1:
            raise HTTPException(409, "当前版本无法审核")
        get_or_404(db, RiskCase, letter.case_id).status = "approved"
        audit(db, "letter_approved", letter_id, {"revision": letter.revision, "actor": "local_test_user"})
        db.commit()
        db.refresh(letter)
        return letter_dict(letter)


@app.post("/api/letters/{letter_id}/simulate-send")
def simulate_send(letter_id: str, data: RevisionPayload):
    """只生成本地回执；同一版本重复调用返回相同回执，无外部发送。"""
    with SessionLocal() as db:
        letter = get_or_404(db, Letter, letter_id)
        ensure_revision(letter, data.expected_revision)
        if letter.status == "simulated_sent" and letter.receipt:
            return letter_dict(letter)
        content_hash = letter_content_hash(letter)
        if letter.status != "approved" or letter.approved_revision != letter.revision or letter.approved_hash != content_hash:
            raise HTTPException(409, "须先审核通过当前函件版本")
        receipt = {"id": f"SIM-{letter.id}-{letter.revision}", "channel": "simulation", "status": "simulated_sent", "created_at": timestamp(now())}
        result = db.execute(update(Letter).where(Letter.id == letter_id, Letter.revision == data.expected_revision, Letter.status == "approved", Letter.approved_hash == content_hash)
                            .values(status="simulated_sent", receipt=receipt))
        if result.rowcount != 1:
            db.rollback()
            db.refresh(letter)
            if letter.status == "simulated_sent":
                return letter_dict(letter)
            raise HTTPException(409, "函件状态已变化")
        get_or_404(db, RiskCase, letter.case_id).status = "simulated_sent"
        audit(db, "letter_simulated_send", letter_id, {"receipt_id": receipt["id"], "no_external_call": True})
        db.commit()
        db.refresh(letter)
        return letter_dict(letter)


@app.get("/api/letters/{letter_id}/download")
def download_letter(letter_id: str):
    with SessionLocal() as db:
        letter = get_or_404(db, Letter, letter_id)
        return Response(letter.content.encode("utf-8-sig"), media_type="text/plain; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="risk-letter-{letter.id}-v{letter.revision}.txt"'})


IMPORT_COLUMNS = ["org_code", "indicator_id", "data_date", "value", "unit", "rank", "population", "rank_scope", "value_encoding"]
ALIASES = {"机构代码": "org_code", "机构号": "org_code", "指标ID": "indicator_id", "指标编号": "indicator_id", "数据日期": "data_date", "指标值": "value", "单位": "unit", "排名": "rank", "全省排名": "rank", "参评数量": "population", "排名范围": "rank_scope", "数值编码": "value_encoding"}


@app.get("/api/import/template.csv")
def import_template():
    output = io.StringIO(newline="")
    csv.writer(output).writerow(IMPORT_COLUMNS)
    return Response(output.getvalue().encode("utf-8-sig"), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="indicator-import-template.csv"'})


def parse_file(content: bytes, filename: str | None) -> tuple[list[str], list[list[Any]]]:
    """CSV/XLSX统一转成表头和原始行；不猜百分比编码、不执行公式。"""
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("gb18030")
            except UnicodeDecodeError:
                raise HTTPException(422, "CSV编码无法识别，请保存为UTF-8")
        rows = list(csv.reader(io.StringIO(text)))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(413, "Excel解压后内容过大")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            if len(book.worksheets) != 1:
                book.close()
                raise HTTPException(422, "请将待导入的指标事实整理为单个工作表")
            rows = []
            for cells in book.worksheets[0].iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    book.close()
                    raise HTTPException(422, "导入不接受公式，请先转换为已核验的数值")
                if any(isinstance(cell.value, (int, float)) and "%" in (cell.number_format or "") for cell in cells):
                    book.close()
                    raise HTTPException(422, "请将百分比单元格转为普通数值（如1.65表示1.65%），并填percent_number编码")
                rows.append([cell.value for cell in cells])
                if len(rows) > 10001:
                    book.close()
                    raise HTTPException(413, "单次导入最多10000条记录")
            book.close()
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, "Excel文件无法读取或格式不受支持")
    else:
        raise HTTPException(422, "仅支持CSV或XLSX文件")
    rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
    if len(rows) < 2:
        raise HTTPException(422, "文件须含表头及至少一条指标事实")
    if len(rows) > 10001:
        raise HTTPException(413, "单次导入最多10000条记录")
    headers = [ALIASES.get(str(x).strip(), str(x).strip().lower()) for x in rows[0]]
    if len(headers) != len(set(headers)):
        raise HTTPException(422, "表头存在重复字段")
    missing = set(IMPORT_COLUMNS) - set(headers)
    if missing:
        raise HTTPException(422, "缺少字段：" + "、".join(sorted(missing)))
    return headers, rows[1:]


def optional_integer(value: Any, label: str, row_number: int) -> int | None:
    """导入排名和参评数量共用的整数解析，空值保持为空。"""
    if value is None or not str(value).strip():
        return None
    try:
        parsed = Decimal(str(value))
    except InvalidOperation:
        raise HTTPException(422, f"第{row_number}行{label}必须为正整数")
    if not parsed.is_finite() or parsed != parsed.to_integral_value() or parsed < 1 or parsed > 100000:
        raise HTTPException(422, f"第{row_number}行{label}必须为正整数")
    return int(parsed)


@app.post("/api/import")
async def import_data(file: UploadFile = File(...), source_kind: str = Form("file")):
    """文件与数据库导入共用此边界；全部行校验后才创建快照。"""
    if source_kind not in {"file", "database"}:
        raise HTTPException(422, "未知的数据来源类型")
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "单次文件大小最多5MB")
    headers, rows = parse_file(content, file.filename)
    with SessionLocal() as db:
        orgs = {o.code: o for o in db.scalars(select(Organization))}
        metrics = {i.id: i for i in db.scalars(select(Indicator))}
        parsed, seen = [], set()
        for position, row in enumerate(rows, 2):
            if len(row) > len(headers) and any(v is not None and str(v).strip() for v in row[len(headers):]):
                raise HTTPException(422, f"第{position}行列数超过表头")
            item = {h: row[i] if i < len(row) else "" for i, h in enumerate(headers)}
            org_code, indicator_id = str(item["org_code"] or "").strip(), str(item["indicator_id"] or "").strip()
            if org_code not in orgs or indicator_id not in metrics:
                raise HTTPException(422, f"第{position}行机构代码或指标ID不在目录中")
            key = (org_code, indicator_id)
            if key in seen:
                raise HTTPException(422, f"第{position}行存在重复机构与指标；一个快照每指标仅接受一个期次")
            seen.add(key)
            raw_date = item["data_date"]
            try:
                day = raw_date.date() if isinstance(raw_date, datetime) else raw_date if isinstance(raw_date, date) else date.fromisoformat(str(raw_date).strip())
            except (ValueError, TypeError):
                raise HTTPException(422, f"第{position}行数据日期须为YYYY-MM-DD")
            metric = metrics[indicator_id]
            unit = str(item["unit"] or "").strip()
            if unit == "CNY":
                unit = "元"
            if unit != metric.unit:
                raise HTTPException(422, f"第{position}行单位应为{metric.unit}，不能自动猜测换算")
            encoding = str(item["value_encoding"] or "").strip()
            expected = "percent_number" if metric.unit == "%" else "raw"
            if encoding != expected:
                raise HTTPException(422, f"第{position}行数值编码应为{expected}")
            raw_value = item["value"]
            try:
                value = None if raw_value is None or not str(raw_value).strip() else Decimal(str(raw_value).strip())
            except InvalidOperation:
                raise HTTPException(422, f"第{position}行指标值须为普通数值或留空")
            if value is not None and (not value.is_finite() or abs(value) >= Decimal("10000000000000000")):
                raise HTTPException(422, f"第{position}行指标值非有限数值或超出存储范围")
            if value is not None and value.as_tuple().exponent < -8:
                raise HTTPException(422, f"第{position}行指标值精度超过8位小数")
            rank = optional_integer(item["rank"], "排名", position)
            population = optional_integer(item["population"], "参评数量", position)
            if rank is not None and population is not None and rank > population:
                raise HTTPException(422, f"第{position}行排名不能大于参评数量")
            parsed.append({"org_id": orgs[org_code].id, "indicator_id": indicator_id, "data_date": day, "value": value, "unit": unit,
                           "rank": rank, "population": population, "rank_scope": str(item["rank_scope"] or "").strip(),
                           "quality_status": "missing" if value is None else "valid"})
        snap = Snapshot(name=f"导入 · {Path(file.filename or 'data').name[:120]}", data_date=max(p["data_date"] for p in parsed),
                        source="database_import" if source_kind == "database" else "file_import", is_demo=False, row_count=len(parsed), status="ready")
        db.add(snap)
        db.flush()
        for item in parsed:
            db.add(Observation(snapshot_id=snap.id, **item))
        audit(db, "snapshot_imported", snap.id, {"row_count": len(parsed), "local_test_only": True})
        db.commit()
        org_levels = {o.id: o.level for o in orgs.values()}
        monitored_orgs = len({p["org_id"] for p in parsed if org_levels[p["org_id"]] == 3})
        warnings = ["导入数据仅用于本地测试；缺失的机构或指标会记录为无法判断。"]
        if monitored_orgs < 82:
            warnings.append(f"本次包含{monitored_orgs}/82家行社的数据，未提供的数据不会当作零值。")
        return {"snapshot_id": snap.id, "row_count": len(parsed), "warnings": warnings}


@app.get("/api/audit")
def audit_logs():
    with SessionLocal() as db:
        return [{"id": a.id, "action": a.action, "entity_id": a.entity_id, "details": a.details, "created_at": timestamp(a.created_at)}
                for a in db.scalars(select(AuditLog).order_by(AuditLog.created_at.desc()).limit(100))]


DIST = ROOT / "frontend" / "dist"
if DIST.is_dir():
    app.mount("/", StaticFiles(directory=str(DIST), html=True), name="frontend")
else:
    @app.get("/")
    def frontend_pending():
        return JSONResponse({"message": "前端尚未构建；请执行前端构建后重启服务。", "api_docs": "/docs"})
